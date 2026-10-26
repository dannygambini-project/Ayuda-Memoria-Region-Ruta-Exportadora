import pandas as pd
import unicodedata
import json
import os
import io
import base64
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def normalizar_texto(texto):
    if pd.isna(texto):
        return ''
    texto = str(texto).strip().lower()
    return ''.join(c for c in unicodedata.normalize('NFD', texto) if unicodedata.category(c) != 'Mn')

def create_regional_docx(region_nombre_oficial, anio_actual, parrafo_1, parrafo_atenciones_header, atenciones_list):
    doc = Document()

    # Título principal
    titulo = f"Ayuda memoria de las atenciones de la Ruta Exportadora en {region_nombre_oficial} {anio_actual}"
    p_title = doc.add_paragraph()
    run_title = p_title.add_run(titulo)
    run_title.bold = True
    run_title.font.size = Pt(14)
    run_title.font.name = 'Arial'
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()  # Espacio

    # Primer párrafo informativo
    p1 = doc.add_paragraph(parrafo_1)
    if p1.runs:
        p1.runs[0].font.name = 'Arial'
        p1.runs[0].font.size = Pt(11)

    # Encabezado de atenciones
    p_head = doc.add_paragraph(parrafo_atenciones_header)
    if p_head.runs:
        p_head.runs[0].bold = True
        p_head.runs[0].font.name = 'Arial'
        p_head.runs[0].font.size = Pt(11)

    # Lista de atenciones
    for item in atenciones_list[:20]:
        evt = item['evento']
        cnt = item['empresas']
        emp_str = "empresa" if cnt == 1 else "empresas"
        p_item = doc.add_paragraph(f"{evt} a {cnt} {emp_str}", style='List Bullet')
        if p_item.runs:
            p_item.runs[0].font.name = 'Arial'
            p_item.runs[0].font.size = Pt(11)

    buf = io.BytesIO()
    doc.save(buf)
    return doc, buf.getvalue()

print("Cargando bases de datos...")
df_rutex = pd.read_excel('base rutex 20.10.2026 - Colab.xlsx')
df_rutex['REGION_NORM'] = df_rutex['REGION'].apply(normalizar_texto)

df_exp = pd.read_excel('EXPORTACIONES 2025.xlsx')
df_exp['RUC_STR'] = df_exp['RUC'].astype(str).str.strip()
df_exp['FOB USD'] = pd.to_numeric(df_exp['FOB USD'], errors='coerce').fillna(0)

df_cat = pd.read_excel('BD EMPRESAS CATEGORIZADAS AL 2025.xlsx')
df_cat['RUC_STR'] = df_cat['RUC'].astype(str).str.strip()
cat_map = dict(zip(df_cat['RUC_STR'], df_cat['CATEGORÍA 2025']))

regions_data = {}
anio_actual = 2026

# Orden alfabético de regiones
unique_regs = df_rutex['REGION'].dropna().unique()
for reg_name in sorted(unique_regs, key=lambda x: normalizar_texto(x)):
    reg_norm = normalizar_texto(reg_name)
    df_rutex_reg = df_rutex[df_rutex['REGION_NORM'] == reg_norm].copy()
    if df_rutex_reg.empty:
        continue
    
    oficial_name = df_rutex_reg['REGION'].iloc[0].strip().title()
    rucs_region = set(df_rutex_reg['RUC'].dropna().astype(str).str.strip())
    total_empresas = len(rucs_region)
    
    eventos_unicos = int(df_rutex_reg['EVENTO'].dropna().nunique())
    df_evt_unicos = df_rutex_reg.drop_duplicates(subset=['EVENTO'])
    evt_lima = int(df_evt_unicos[df_evt_unicos['REGION_EVENTO'].astype(str).str.strip().str.upper() == 'LIMA'].shape[0])
    evt_regiones = eventos_unicos - evt_lima
    
    # Sectores
    df_sector = df_rutex_reg.drop_duplicates(subset=['RUC', 'SECTOR'])
    sectores_counts = df_sector['SECTOR'].dropna().str.strip().value_counts()
    sectores_list = []
    sectores_str_list = []
    for s_nom, s_cant in sectores_counts.items():
        pct = round((s_cant / total_empresas) * 100) if total_empresas > 0 else 0
        sectores_list.append({'sector': s_nom, 'cantidad': int(s_cant), 'porcentaje': pct})
    for item in sectores_list[:3]:
        sectores_str_list.append(f"{item['sector']} con {item['porcentaje']}%")
    
    if len(sectores_str_list) > 1:
        sectores_texto = ', '.join(sectores_str_list[:-1]) + ' e ' + sectores_str_list[-1]
    elif len(sectores_str_list) == 1:
        sectores_texto = sectores_str_list[0]
    else:
        sectores_texto = 'sin sectores registrados'
        
    # Exportaciones
    df_exp_filtrada = df_exp[df_exp['RUC_STR'].isin(rucs_region)]
    df_exp_validas = df_exp_filtrada[df_exp_filtrada['FOB USD'] > 0]
    rucs_exportadores = set(df_exp_validas['RUC_STR'])
    cant_exportadoras = len(rucs_exportadores)
    cant_mercados = int(df_exp_validas['MERCADO'].dropna().nunique())
    cant_productos = int(df_exp_validas['PARTIDA_DESCRIP_LARGA'].dropna().nunique())
    monto_fob_total = float(df_exp_validas['FOB USD'].sum())
    monto_fob_mill = monto_fob_total / 1_000_000
    
    # Top Productos
    top_productos_series = df_exp_validas.groupby('PARTIDA_DESCRIP_LARGA')['FOB USD'].sum().sort_values(ascending=False)
    prod_list = []
    prod_str_list = []
    for p_nom, p_monto in top_productos_series.head(10).items():
        m_mill = p_monto / 1_000_000
        p_str = str(p_nom).strip()
        display_nom = p_str.split('-', 1)[1].strip() if ('-' in p_str and p_str.split('-', 1)[0].strip().isdigit()) else p_str
        prod_list.append({
            'partida_completa': p_str,
            'producto': display_nom,
            'fob_usd': float(p_monto),
            'fob_millones': round(m_mill, 1),
            'participacion': round((p_monto / monto_fob_total * 100), 1) if monto_fob_total > 0 else 0
        })
    for item in prod_list[:4]:
        prod_str_list.append(f"{item['producto'].lower()} (USD {item['fob_millones']:,.1f} millones)")
    if prod_str_list:
        prod_texto = ', '.join(prod_str_list[:-1]) + ' y ' + prod_str_list[-1] if len(prod_str_list) > 1 else prod_str_list[0]
    else:
        prod_texto = 'sin productos destacados'
        
    # Top Mercados
    top_mercados_series = df_exp_validas.groupby('MERCADO')['FOB USD'].sum().sort_values(ascending=False)
    merc_list = []
    merc_str_list = []
    for m_nom, m_monto in top_mercados_series.head(10).items():
        m_mill = m_monto / 1_000_000
        merc_list.append({
            'mercado': str(m_nom).strip().title(),
            'fob_usd': float(m_monto),
            'fob_millones': round(m_mill, 1),
            'participacion': round((m_monto / monto_fob_total * 100), 1) if monto_fob_total > 0 else 0
        })
    for item in merc_list[:3]:
        merc_str_list.append(f"{item['mercado']} (USD {item['fob_millones']:,.1f} millones)")
    if merc_str_list:
        merc_texto = ', '.join(merc_str_list[:-1]) + ' y ' + merc_str_list[-1] if len(merc_list) > 1 else merc_str_list[0]
    else:
        merc_texto = 'sin mercados destacados'
        
    # Top Atenciones
    top_atenciones_series = df_rutex_reg['EVENTO'].dropna().astype(str).str.strip().value_counts()
    atenciones_list = []
    for evt_nom, evt_cnt in top_atenciones_series.head(20).items():
        atenciones_list.append({
            'evento': evt_nom,
            'empresas': int(evt_cnt)
        })
        
    # Categorías 2025
    cats_count = {'Esporádico': 0, 'Regular': 0, 'Nuevas': 0, 'No exportó': 0, 'Sin categoría': 0}
    for r in rucs_region:
        c = cat_map.get(r, 'Sin categoría')
        cats_count[c] = cats_count.get(c, 0) + 1
        
    # Párrafo generado idéntico al notebook
    monto_fob_display = f"{round(monto_fob_mill):,}"
    parrafo_1 = (
        f"Durante el {anio_actual} se vienen atendiendo a {total_empresas} empresas de la región de {oficial_name}, "
        f"a través de {eventos_unicos} actividades, {evt_lima} en Lima y {evt_regiones} en regiones. "
        f"Los principales sectores son {sectores_texto}. De las {total_empresas} empresas atendidas, "
        f"{cant_exportadoras} exportaron a {cant_mercados} mercados, {cant_productos} productos un monto de USD {monto_fob_display} millones, "
        f"siendo sus principales productos {prod_texto}. Sus principales mercados atendidos son {merc_texto}."
    )
    
    header_atenciones = f"Las principales atenciones recibidas en el {anio_actual} son:"
    doc_obj, docx_bytes = create_regional_docx(oficial_name, anio_actual, parrafo_1, header_atenciones, atenciones_list)
    docx_b64 = base64.b64encode(docx_bytes).decode('utf-8')

    os.makedirs('documentos_word', exist_ok=True)
    doc_path = os.path.join('documentos_word', f"Ayuda memoria de las atenciones de la Ruta Exportadora en {oficial_name} {anio_actual}.docx")
    doc_obj.save(doc_path)

    regions_data[reg_norm] = {
        'id': reg_norm,
        'nombre': oficial_name,
        'anio': anio_actual,
        'total_empresas': total_empresas,
        'eventos_unicos': eventos_unicos,
        'evt_lima': evt_lima,
        'evt_regiones': evt_regiones,
        'sectores': sectores_list,
        'sectores_texto': sectores_texto,
        'cant_exportadoras': cant_exportadoras,
        'cant_mercados': cant_mercados,
        'cant_productos': cant_productos,
        'monto_fob_total': monto_fob_total,
        'monto_fob_millones': round(monto_fob_mill, 1),
        'monto_fob_display': monto_fob_display,
        'productos': prod_list,
        'productos_texto': prod_texto,
        'mercados': merc_list,
        'mercados_texto': merc_texto,
        'atenciones': atenciones_list,
        'categorias': cats_count,
        'parrafo_narrativa': parrafo_1,
        'parrafo_atenciones_header': header_atenciones,
        'titulo_documento': f"Ayuda memoria de las atenciones de la Ruta Exportadora en {oficial_name} {anio_actual}",
        'docx_base64': docx_b64
    }

# Estadísticas globales del Perú (para resumen nacional comparativo)
total_empresas_peru = len(set(df_rutex['RUC'].dropna().astype(str).str.strip()))
eventos_totales_peru = int(df_rutex['EVENTO'].dropna().nunique())

output_obj = {
    'anio': anio_actual,
    'total_regiones': len(regions_data),
    'regiones': regions_data,
    'regiones_lista': [{'id': k, 'nombre': v['nombre'], 'empresas': v['total_empresas'], 'fob_millones': v['monto_fob_millones']} for k, v in sorted(regions_data.items(), key=lambda x: x[1]['nombre'])]
}

with open('regiones_data.json', 'w', encoding='utf-8') as f:
    json.dump(output_obj, f, ensure_ascii=False, indent=2)

print(f"Éxito: {len(regions_data)} regiones procesadas y guardadas en regiones_data.json")
